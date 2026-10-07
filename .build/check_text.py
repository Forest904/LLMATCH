from pptx import Presentation
from PIL import ImageFont

p = Presentation('output/LLMATCH_presentazione_10_slide_finale.pptx')
regular = 'C:/Windows/Fonts/segoeui.ttf'
bold = 'C:/Windows/Fonts/segoeuib.ttf'
bad = []
for si, slide in enumerate(p.slides, 1):
    for shape in slide.shapes:
        if not shape.has_text_frame or not shape.text.strip():
            continue
        lines = 0
        max_size = 0
        for para in shape.text_frame.paragraphs:
            if not para.text.strip():
                lines += 1
                continue
            run = para.runs[0]
            size = run.font.size.pt if run.font.size else 18
            font = ImageFont.truetype(bold if run.font.bold else regular, round(size*96/72))
            max_size = max(max_size, size)
            width = shape.width / 914400 * 96 - 6
            current = ''
            for word in para.text.split():
                next_line = (current + ' ' + word).strip()
                if font.getlength(next_line) > width and current:
                    lines += 1
                    current = word
                else:
                    current = next_line
            if current:
                lines += 1
        needed = lines * max_size * 96/72 * 1.12
        available = shape.height / 914400 * 96
        if needed > available + 6:
            bad.append((si, shape.text[:50], round(needed), round(available)))
print('potential_overflow:', bad)
