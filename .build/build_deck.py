from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / 'Assets'
OUT = ROOT / 'output'
OUT.mkdir(exist_ok=True)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
NAVY = RGBColor(20, 42, 60)
TEAL = RGBColor(16, 123, 131)
GRAY = RGBColor(80, 93, 103)
LIGHT = RGBColor(224, 232, 235)
BG = RGBColor(250, 251, 250)

def txt(slide, text, x, y, w, h, size=20, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Segoe UI'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
        p.space_after = Pt(7)
        if align is not None: p.alignment = align
    return box

def base(title, n, source=None):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = BG
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.085))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = TEAL; stripe.line.fill.background()
    txt(slide, title, .72, .46, 11.95, .86, 30, NAVY, True)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.72), Inches(1.37), Inches(11.9), Inches(.018))
    line.fill.solid(); line.fill.fore_color.rgb = LIGHT; line.line.fill.background()
    if source:
        txt(slide, source, .72, 7.12, 11.9, .22, 9, GRAY)
    txt(slide, f'{n:02d}', 12.25, 7.08, .42, .26, 10, GRAY, align=PP_ALIGN.RIGHT)
    return slide

def pic(slide, name, x, y, w, h, crop=None):
    with Image.open(ASSETS/name) as im:
        iw, ih = im.size
    l, t, r, b = crop or (0, 0, 0, 0)
    aspect = (iw * (1-l-r)) / (ih * (1-t-b))
    if w / h > aspect:
        fw, fh = h * aspect, h
    else:
        fw, fh = w, w / aspect
    p = slide.shapes.add_picture(str(ASSETS/name), Inches(x+(w-fw)/2), Inches(y+(h-fh)/2), width=Inches(fw), height=Inches(fh))
    if crop:
        p.crop_left, p.crop_top, p.crop_right, p.crop_bottom = crop
    return p

def rule(slide, x, y, w):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(.025))
    s.fill.solid(); s.fill.fore_color.rgb = LIGHT; s.line.fill.background()

# 1 — title
s = base('LLMATCH', 1, 'Wang et al. · arXiv:2507.10897v1 (2025)')
txt(s, 'Schema matching multi-tabella con Large Language Models', .78, 1.75, 11.5, 1.45, 35, NAVY, True)
txt(s, 'Un framework modulare che usa Rollup e Drilldown per allineare schemi relazionali complessi.', .8, 3.53, 10.8, 1.02, 23, TEAL)
txt(s, 'Sha Wang · Yuchen Li · Hanhua Xiao · Bing Tian Dai · Roy Ka-Wei Lee · Yanfei Dong · Lambert Deng', .8, 5.18, 11.7, .55, 15, GRAY)
txt(s, 'Problema  ·  Metodo  ·  Evidenza sperimentale', .8, 5.93, 10.8, .47, 18, GRAY)

# 2 — problem
s = base('Il significato di una colonna dipende dallo schema', 2, 'Fonte: Wang et al., Fig. 1 e §1–3')
pic(s, 'Fig1.png', .75, 1.76, 7.35, 4.81)
txt(s, 'Person.address ≠ ShopInfo.address', 8.45, 2.04, 4.1, .7, 21, TEAL, True)
txt(s, 'Nomi simili possono indicare entità diverse.\n\nPK/FK e struttura delle tabelle chiariscono il contesto.\n\nI mapping reali possono essere n:m.', 8.45, 2.95, 4.15, 3.3, 18, NAVY)

# 3 — architecture
s = base('Tre stadi per il matching complesso', 3, 'Fonte: Wang et al., Fig. 2 e §4')
txt(s, 'Schema Preparation organizza nomi, descrizioni e relazioni PK/FK; Table Selection restringe i candidati; Column Matching produce i mapping.', .76, 1.62, 11.78, 1.1, 20, NAVY)
pic(s, 'Fig2.png', .76, 2.94, 11.78, 3.58, crop=(0, 0, 0, .155))
txt(s, 'Il focus pairwise dei metodi precedenti perde parte del contesto multi-tabella.', .78, 6.63, 11.8, .35, 16, TEAL)

# 4 — rollup
s = base('Rollup: un’astrazione tracciabile', 4, 'Fonte: Wang et al., Fig. 3 e §4.1')
txt(s, 'L’LLM raggruppa colonne semanticamente correlate in un alias temporaneo.', .77, 1.62, 11.65, .55, 20, NAVY)
pic(s, 'Fig3.png', .75, 2.33, 11.85, 4.05, crop=(0, 0, 0, .385))
txt(s, 'currency, currency_code e currency_id → currency_meta_data; gli originali restano recuperabili.', .78, 6.52, 11.7, .42, 16, TEAL)

# 5 — table selection
s = base('Table Selection: quante tabelle arrivano al matcher?', 5, 'Fonte: Wang et al., §4.2')
rows = [
    ('None', 'Tutte le tabelle target in un unico contesto', 'Contesto esteso'),
    ('Nested Join', 'Un confronto per ogni coppia source–target', 'Contesto minimo'),
    ('Vector Similarity', 'Top-k secondo similarità degli embedding', 'Filtro lessicale/semantico'),
    ('LLM Selection', 'L’LLM sceglie dalle descrizioni delle tabelle', 'Filtro contestuale'),
]
txt(s, 'Strategia', .78, 1.76, 2.35, .36, 16, TEAL, True)
txt(s, 'Criterio', 3.18, 1.76, 6.17, .36, 16, TEAL, True)
txt(s, 'Effetto sul prompt', 9.52, 1.76, 2.8, .36, 16, TEAL, True)
rule(s, .78, 2.2, 11.6)
for i, (a,b,c) in enumerate(rows):
    y = 2.48 + i*.9
    txt(s, a, .78, y, 2.37, .72, 18, NAVY, True)
    txt(s, b, 3.18, y, 6.16, .72, 18, NAVY)
    txt(s, c, 9.52, y, 2.82, .72, 16, GRAY)
    rule(s, .78, y+.73, 11.6)
txt(s, 'La scelta bilancia dimensione del contesto e rischio di escludere match validi.', .78, 6.43, 11.55, .55, 19, TEAL)

# 6 — drilldown
s = base('Drilldown: dal match grossolano alle colonne finali', 6, 'Fonte: Wang et al., Fig. 3 e §4.3')
txt(s, 'Il matching iniziale può coinvolgere un alias creato da Rollup. Drilldown riapre solo quel gruppo e rivaluta le colonne originali.', .77, 1.62, 11.7, .95, 19, NAVY)
pic(s, 'Fig3.png', .75, 2.72, 11.85, 3.78, crop=(0, .39, 0, 0))
txt(s, 'bal_currency → currency e currency_code; currency_id non è necessariamente incluso.', .78, 6.62, 11.7, .33, 16, TEAL)

# 7 — dataset
s = base('SchemaNet mette alla prova schemi relazionali complessi', 7, 'Fonte: Wang et al., Tab. 1 e §5–6.2')
txt(s, '7 coppie source–target  ·  healthcare, finance, entertainment  ·  F1 come metrica', .78, 1.61, 11.65, .55, 20, TEAL, True)
txt(s, 'mimic_iii → omop: 26 tabelle e 324 colonne → 39 tabelle e 432 colonne; 189 mapping.', .78, 2.28, 11.7, .6, 19, NAVY)
pic(s, 'Tab1.png', .77, 3.02, 11.65, 3.12, crop=(0, .20, 0, 0))
txt(s, 'Baseline: Coma, SF, Cupid, Unicorn, ReMatch. Modelli: GPT-3.5-turbo e GPT-4o-mini; SBERT per embedding; nessun fine-tuning.', .78, 6.35, 11.63, .6, 15, GRAY)

# 8 — end-to-end
s = base('Il vantaggio emerge nei task complessi', 8, 'Fonte: Wang et al., Tab. 2 e §6.4. Per mimic_iii-omop si usa il valore 0,33 della tabella; il testo riporta 0,4.')
txt(s, 'mimic_iii → omop · F1', .78, 1.65, 11.7, .5, 19, TEAL, True)
vals = [('Coma', '0,04'), ('SF', '0,09'), ('ReMatch GPT-4o-mini', '0,20'), ('LLMATCH GPT-4o-mini', '0,33')]
for i, (label, value) in enumerate(vals):
    x = .8 + i*3.02
    txt(s, value, x, 2.31, 2.75, .94, 34, TEAL if i==3 else NAVY, True, PP_ALIGN.CENTER)
    txt(s, label, x, 3.25, 2.75, .76, 15, NAVY, align=PP_ALIGN.CENTER)
pic(s, 'Tab2.png', .8, 4.32, 11.73, 2.33, crop=(0, .20, 0, .08))
txt(s, 'bank1-bank2 è un’eccezione: il vocabolario condiviso favorisce anche i matcher lessicali.', .8, 6.72, 11.6, .35, 15, GRAY)

# 9 — components
s = base('I moduli incidono soprattutto sui casi complessi', 9, 'Fonte: Wang et al., Tab. 3, Fig. 5–6 e §6.5–6.7')
txt(s, 'Table Selection · Tab. 3', .78, 1.63, 5.5, .42, 19, TEAL, True)
txt(s, 'mimic_iii–omop', .78, 2.2, 5.45, .4, 18, NAVY, True)
txt(s, 'LLM: F1 0,33 con 6,9 tabelle\nVector Similarity top 10: F1 0,25 con 10 tabelle', .78, 2.67, 5.45, 1.05, 17, NAVY)
txt(s, 'synthea–omop', .78, 4.16, 5.45, .4, 18, NAVY, True)
txt(s, 'LLM: F1 0,37 con 4,7 tabelle\nVector Similarity top 10: F1 0,29 con 10 tabelle', .78, 4.63, 5.45, 1.05, 17, NAVY)
txt(s, 'Rollup / Drilldown · Fig. 6', 6.72, 1.63, 5.45, .42, 19, TEAL, True)
txt(s, 'Il guadagno cresce sugli schemi più frammentati e relazionali.', 6.72, 2.1, 5.45, .92, 17, NAVY)
pic(s, 'Fig6.png', 6.86, 3.24, 5.65, 3.0)
txt(s, 'Descrizioni e relazioni PK/FK migliorano ulteriormente l’F1 (Fig. 5).', .78, 6.45, 11.55, .5, 17, GRAY)

# 10 — limits
s = base('Evidenza promettente, generalizzazione ancora aperta', 10, 'Fonte: Wang et al., Fig. 7–8 e §6.8–7')
txt(s, 'Contesto', .78, 1.67, 3.6, .42, 19, TEAL, True)
txt(s, 'Finestre troppo piccole frammentano il task e riducono l’F1.', .78, 2.12, 3.75, 1.2, 18, NAVY)
txt(s, 'Assistenza umana', 4.85, 1.67, 3.6, .42, 19, TEAL, True)
txt(s, 'F1 medio 0,14 → 0,58 in task da 2 minuti, su 4 partecipanti.', 4.85, 2.12, 3.75, 1.2, 18, NAVY)
txt(s, 'Questioni aperte', 8.93, 1.67, 3.6, .42, 19, TEAL, True)
txt(s, 'Non-determinismo, campione ridotto e domini non rappresentati.', 8.93, 2.12, 3.55, 1.2, 18, NAVY)
pic(s, 'Fig7.png', .78, 3.68, 5.66, 1.15)
pic(s, 'Fig8.png', 7.55, 3.55, 4.42, 2.89)
txt(s, 'LLMATCH convince sui task complessi; riproducibilità e uso pratico richiedono verifiche più ampie.', .78, 6.52, 11.7, .52, 17, TEAL, True)

path = OUT / 'LLMATCH_presentazione_10_slide_finale.pptx'
prs.save(path)
print(path)
